import sys
try:
    from docx import Document
except ImportError:
    print("python-docx is not installed.")
    sys.exit(1)

def check_docx(doc_path):
    doc = Document(doc_path)
    full_text = []
    
    # Extract from paragraphs
    for p in doc.paragraphs:
        full_text.append(p.text)
            
    # Extract from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full_text.append(p.text)
                    
    text = "\n".join(full_text)
    
    # Check for red flags
    red_flags = ["YOLOv8", "Java backend", "92.4", "94%"]
    
    for flag in red_flags:
        if flag in text:
            print(f"RED FLAG FOUND: '{flag}'")
            
    print("--- DOCX Content Snippet ---")
    print(text[:2000])

if __name__ == "__main__":
    docx_file = r"Z:\major\check_docx\KrushikaDhara_IEEE_PUBLISH_READY.docx"
    check_docx(docx_file)
