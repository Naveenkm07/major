import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation(out_path):
    prs = Presentation()
    
    # ---------------------------------------------------------
    # Slide 1: Title Slide
    # ---------------------------------------------------------
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "KrushikaDhara: Novelty & Uniqueness"
    subtitle.text = "Solving the Integration Gap in Indian Agriculture\nThrough Offline-First Edge AI & Bluetooth Mesh Networks"
    
    # ---------------------------------------------------------
    # Slide 2: The Integration Gap
    # ---------------------------------------------------------
    layout = prs.slide_layouts[5] # blank with title
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "1. Solving the 'Integration Gap'"
    
    # Add text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "Current Limitations in Research:"
    p.font.bold = True
    p.font.size = Pt(20)
    
    bullets = [
        "Most agri-tech research produces isolated models (only disease detection, or only yield).",
        "Farmers suffer from 'app fatigue' trying to use 10 different tools."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(16)
        
    p = tf.add_paragraph()
    p.text = "\nKrushikaDhara's Unique Approach:"
    p.font.bold = True
    p.font.size = Pt(20)
    
    bullets2 = [
        "Unifies 9 advanced modules into a single, cohesive workflow.",
        "Combines Disease Detection, Yield Estimation, RAG-Advisory, P2P Rental, and Market Prices.",
        "Built entirely on free-tier infrastructure, ensuring zero recurring cost for NGOs."
    ]
    for b in bullets2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(16)
        
    # Add Image
    img_path = r"C:\Users\indar\.gemini\antigravity-ide\brain\a51ad55d-cc05-4a05-b053-b7572930ee7a\farmer_app_ui_1790370010020.jpg"
    try:
        slide.shapes.add_picture(img_path, Inches(5.5), Inches(1.5), width=Inches(4))
    except Exception as e:
        print("Image 1 failed to load:", e)
        
    # ---------------------------------------------------------
    # Slide 3: Multi-Tiered AI
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "2. Multi-Tiered AI (Edge + Aerial Fusion)"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "Tier 1: Micro-Level (Edge AI)"
    p.font.bold = True
    p.font.size = Pt(20)
    p = tf.add_paragraph()
    p.text = "• INT8-quantised YOLOv8 runs directly on the farmer's smartphone."
    p.font.size = Pt(16)
    p = tf.add_paragraph()
    p.text = "• Zero internet required, median latency of just 178.4ms."
    p.font.size = Pt(16)
    
    p = tf.add_paragraph()
    p.text = "\nTier 2: Macro-Level (Aerial Fusion)"
    p.font.bold = True
    p.font.size = Pt(20)
    p = tf.add_paragraph()
    p.text = "• Fuses high-res Drone imagery with Sentinel-2 satellite NDVI data."
    p.font.size = Pt(16)
    p = tf.add_paragraph()
    p.text = "• Generates highly accurate per-acre yield estimates and pest hotspot maps."
    p.font.size = Pt(16)

    img_path = r"C:\Users\indar\.gemini\antigravity-ide\brain\a51ad55d-cc05-4a05-b053-b7572930ee7a\drone_field_1790370023751.jpg"
    try:
        slide.shapes.add_picture(img_path, Inches(5.5), Inches(1.5), width=Inches(4))
    except Exception:
        pass
        
    # ---------------------------------------------------------
    # Slide 4: Offline P2P Equipment Rental
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "3. Bluetooth Mesh Equipment Rental"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "The Connectivity Challenge:"
    p.font.bold = True
    p.font.size = Pt(20)
    p = tf.add_paragraph()
    p.text = "• Traditional rental apps break when a farmer loses 4G signal in the field."
    p.font.size = Pt(16)
    
    p = tf.add_paragraph()
    p.text = "\nThe Offline Fallback Mechanism:"
    p.font.bold = True
    p.font.size = Pt(20)
    
    b3 = [
        "Uses Firebase Geohashing when online.",
        "Fails over to a Bluetooth/WiFi-Direct Mesh when offline.",
        "Farmers broadcast equipment listings directly to nearby phones without a cell tower.",
        "Allows true Peer-to-Peer (P2P) resource pooling, completely bypassing infrastructure failures."
    ]
    for b in b3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(16)

    img_path = r"C:\Users\indar\.gemini\antigravity-ide\brain\a51ad55d-cc05-4a05-b053-b7572930ee7a\tractor_bluetooth_1790370034895.jpg"
    try:
        slide.shapes.add_picture(img_path, Inches(5.5), Inches(1.5), width=Inches(4))
    except Exception:
        pass
        
    # ---------------------------------------------------------
    # Slide 5: Vernacular Voice LLM
    # ---------------------------------------------------------
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "4. Vernacular RAG Advisory System"
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = "Eliminating the Literacy Barrier:"
    p.font.bold = True
    p.font.size = Pt(20)
    p = tf.add_paragraph()
    p.text = "• Text-based LLMs exclude farmers who cannot read or type in English/Hindi."
    p.font.size = Pt(16)
    
    p = tf.add_paragraph()
    p.text = "\nBhashini Voice Pipeline:"
    p.font.bold = True
    p.font.size = Pt(20)
    
    b4 = [
        "Integrates Bhashini (ASR + NMT + TTS) for native Kannada voice interactions.",
        "Farmers speak their questions, and the system replies entirely via voice."
    ]
    for b in b4:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(16)
        
    p = tf.add_paragraph()
    p.text = "\nHallucination-Free Responses:"
    p.font.bold = True
    p.font.size = Pt(20)
    p = tf.add_paragraph()
    p.text = "• Uses Retrieval-Augmented Generation (RAG) strictly grounded in official Karnataka welfare scheme documents."
    p.font.size = Pt(16)

    img_path = r"C:\Users\indar\.gemini\antigravity-ide\brain\a51ad55d-cc05-4a05-b053-b7572930ee7a\voice_ai_1790370047858.jpg"
    try:
        slide.shapes.add_picture(img_path, Inches(5.5), Inches(1.5), width=Inches(4))
    except Exception:
        pass

    prs.save(out_path)
    print(f"Presentation successfully created at {out_path}")

if __name__ == "__main__":
    out_path = r"Z:\major\KrushikaDhara_Uniqueness_Pitch.pptx"
    create_presentation(out_path)
