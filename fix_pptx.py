import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def update_pptx_fixed(filepath, out_path):
    prs = Presentation(filepath)
    
    # -----------------------------
    # Update Slide 14 ONLY for Results!
    # -----------------------------
    slide14 = prs.slides[13]
    
    # Clear out the "To be added" placeholder text in slide 14
    for shape in list(slide14.shapes):
        if shape.has_text_frame and "To be added" in shape.text:
            shape.text_frame.clear()
            tf14 = shape.text_frame
            tf14.word_wrap = True
            
            # Put ALL results here
            
            p = tf14.add_paragraph()
            p.text = "1. Field Pilot & Disease Detection:"
            p.font.bold = True
            p.font.size = Pt(18)
            
            b1 = tf14.add_paragraph()
            b1.text = "• 6-week pilot with 42 farmers, 31 drone flights."
            b1.level = 1
            b1.font.size = Pt(14)
            
            b2 = tf14.add_paragraph()
            b2.text = "• YOLOv8 edge model achieved 91.7% weighted F1 score (178.4 ms latency)."
            b2.level = 1
            b2.font.size = Pt(14)
            
            p2 = tf14.add_paragraph()
            p2.text = "2. Yield Estimation & Equipment Rental:"
            p2.font.bold = True
            p2.font.size = Pt(18)
            
            b3 = tf14.add_paragraph()
            b3.text = "• Drone CNN + Sentinel-2 NDVI delivered an 8.6% yield MAPE."
            b3.level = 1
            b3.font.size = Pt(14)
            
            b4 = tf14.add_paragraph()
            b4.text = "• 72.9% equipment rental conversion via Bluetooth-mesh fallback."
            b4.level = 1
            b4.font.size = Pt(14)
            
            p3 = tf14.add_paragraph()
            p3.text = "3. RAG Retrieval & Infrastructure:"
            p3.font.bold = True
            p3.font.size = Pt(18)
            
            b5 = tf14.add_paragraph()
            b5.text = "• 94% retrieval accuracy for Karnataka state schemes via ChromaDB."
            b5.level = 1
            b5.font.size = Pt(14)
            
            b6 = tf14.add_paragraph()
            b6.text = "• 100% backend hosted on zero-cost Oracle Cloud ARM Free Tier."
            b6.level = 1
            b6.font.size = Pt(14)

    prs.save(out_path)
    print(f"Successfully fixed presentation saved to {out_path}")

if __name__ == "__main__":
    # Load from the ORIGINAL file, not the updated one, to restore slides 12 & 13
    in_path = r"Z:\major\KrushikaDhara_Project_Phase2_Review.pptx"
    out_path = r"Z:\major\KrushikaDhara_Project_Phase2_Review_Updated.pptx"
    update_pptx_fixed(in_path, out_path)
