import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def update_pptx(filepath, out_path):
    prs = Presentation(filepath)
    
    # -----------------------------
    # Update Slide 12: Field Pilot & Disease Detection
    # -----------------------------
    slide12 = prs.slides[11]
    title12 = slide12.shapes.title
    if title12 is None:
        # Create title if missing
        title12 = slide12.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
    title12.text = "Results: Field Pilot & Disease Detection"
    
    content12 = slide12.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf12 = content12.text_frame
    tf12.word_wrap = True
    
    p = tf12.add_paragraph()
    p.text = "Field Pilot Overview:"
    p.font.bold = True
    p.font.size = Pt(24)
    
    bullets12 = [
        "Duration: 6 weeks across Chitradurga and Kolar districts.",
        "Participation: 42 farmers actively engaged.",
        "Activity: 31 drone flights conducted, 96 equipment listings created.",
    ]
    for b in bullets12:
        p = tf12.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(20)
        
    p = tf12.add_paragraph()
    p.text = "Disease Detection Performance (Table III):"
    p.font.bold = True
    p.font.size = Pt(24)
    
    bullets12_b = [
        "INT8-quantised YOLOv8 achieved a 91.7% weighted F1 score.",
        "Median on-device inference latency: 178.4 ms.",
        "Zero server dependency, proving the efficacy of offline-first design."
    ]
    for b in bullets12_b:
        p = tf12.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(20)

    # -----------------------------
    # Update Slide 13: Yield Estimation & Rental
    # -----------------------------
    slide13 = prs.slides[12]
    title13 = slide13.shapes.title
    if title13 is None:
        title13 = slide13.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
    title13.text = "Results: Yield Estimation & Equipment Rental"
    
    content13 = slide13.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf13 = content13.text_frame
    tf13.word_wrap = True
    
    p = tf13.add_paragraph()
    p.text = "Drone-Based Yield & Hotspot Accuracy (Table IV):"
    p.font.bold = True
    p.font.size = Pt(24)
    
    bullets13 = [
        "CNN + Sentinel-2 NDVI fusion delivered an 8.6% Mean Absolute Percentage Error (MAPE).",
        "Successfully identified overlapping pest stress zones and nutrient deficiency.",
    ]
    for b in bullets13:
        p = tf13.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(20)
        
    p = tf13.add_paragraph()
    p.text = "Equipment Marketplace Conversion:"
    p.font.bold = True
    p.font.size = Pt(24)
    
    bullets13_b = [
        "Achieved a 72.9% equipment rental conversion rate during the pilot.",
        "Bluetooth-mesh geohash fallback enabled peer discovery even during 4G outages."
    ]
    for b in bullets13_b:
        p = tf13.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(20)

    # -----------------------------
    # Update Slide 14: Infrastructure & RAG
    # -----------------------------
    slide14 = prs.slides[13]
    # We will clear out the placeholder text in slide 14 first
    for shape in list(slide14.shapes):
        if shape.has_text_frame and "To be added" in shape.text:
            shape.text_frame.clear()
            tf14 = shape.text_frame
            
            p = tf14.add_paragraph()
            p.text = "Scheme RAG Retrieval Evaluation (Table V):"
            p.font.bold = True
            p.font.size = Pt(24)
            
            bullets14 = [
                "Achieved 94% retrieval accuracy for Karnataka state schemes.",
                "Sub-40 ms cosine search via ChromaDB vector store."
            ]
            for b in bullets14:
                p = tf14.add_paragraph()
                p.text = b
                p.level = 1
                p.font.size = Pt(20)
                
            p = tf14.add_paragraph()
            p.text = "Infrastructure / Cost Breakdown (Table II):"
            p.font.bold = True
            p.font.size = Pt(24)
            
            bullets14_b = [
                "100% of the backend runs on Oracle Cloud's Always-Free ARM tier.",
                "Zero recurring infrastructure costs to serve the entire 42-farmer pilot."
            ]
            for b in bullets14_b:
                p = tf14.add_paragraph()
                p.text = b
                p.level = 1
                p.font.size = Pt(20)

    prs.save(out_path)
    print(f"Successfully updated presentation saved to {out_path}")

if __name__ == "__main__":
    in_path = r"Z:\major\KrushikaDhara_Project_Phase2_Review.pptx"
    out_path = r"Z:\major\KrushikaDhara_Project_Phase2_Review_Updated.pptx"
    update_pptx(in_path, out_path)
