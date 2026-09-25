import collections 
import collections.abc
import sys
from pptx import Presentation

def extract_text_from_pptx(filepath, out_path):
    try:
        prs = Presentation(filepath)
        with open(out_path, 'w', encoding='utf-8') as f:
            for i, slide in enumerate(prs.slides):
                f.write(f"--- Slide {i+1} ---\n")
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        f.write(shape.text + "\n")
                f.write("\n")
    except Exception as e:
        print(f"Error reading pptx: {e}")

if __name__ == "__main__":
    filepath = r"Z:\major\KrushikaDhara_Project_Phase2_Review.pptx"
    out_path = r"Z:\major\extracted_pptx_clean.txt"
    extract_text_from_pptx(filepath, out_path)
    print("Done")
