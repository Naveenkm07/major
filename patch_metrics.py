import sys
try:
    from docx import Document
except ImportError:
    print("python-docx is not installed.")
    sys.exit(1)

def replace_text_in_docx(doc_path, replacements):
    doc = Document(doc_path)
    replaced_count = 0
    
    for old_text, new_text in replacements.items():
        # Replace in paragraphs
        for p in doc.paragraphs:
            if old_text in p.text:
                p.text = p.text.replace(old_text, new_text)
                replaced_count += 1
                
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if old_text in p.text:
                            p.text = p.text.replace(old_text, new_text)
                            replaced_count += 1
                            
    if replaced_count > 0:
        doc.save(doc_path)
        print(f"Success! Made {replaced_count} total replacements.")
    else:
        print("Notice: No replacements made.")

if __name__ == "__main__":
    docx_file = r"Z:\major\check_docx\KrushikaDhara_IEEE_PUBLISH_READY.docx"
    
    replacements = {
        "YOLOv8 (INT8 Quantized)": "MobileNetV3-Small (INT8 Quantized)",
        "YOLOv8 object detection model to enable precise multi-instance pest localization": "MobileNetV3-Small classification model for lightweight pest detection",
        "quantized YOLOv8 .tflite model": "quantized MobileNetV3 .tflite model",
        "YOLOv8 TFLite interpreter": "MobileNetV3 TFLite interpreter",
        "92.4": "69.4",
        "90.1": "68.2",
        "0.912": "0.688",
        "91.7": "70.1",
        "93.2": "71.5",
        "0.924": "0.708",
        "89.3": "67.4",
        "88.6": "66.9",
        "0.889": "0.671",
        "94.1": "71.8",
        "91.8": "70.4",
        "0.929": "0.711",
        "93.5": "71.1",
        "92.9": "69.8",
        "0.932": "0.704",
        "92.2": "69.9",
        "91.3": "69.4",
        "0.917": "0.696",
        "0.917": "0.696"
    }
    
    replace_text_in_docx(docx_file, replacements)
